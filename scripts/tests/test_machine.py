"""End-to-end dry run of the state machine with mocked agents and mocked notebook answers.

It verifies the plugin's own logic (phases, ledger rules, register, rendering). It does NOT verify anything about
NotebookLM, Legal Data Hunter or the agents' prompts: the agents here are plain functions.
"""

import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
TMP = tempfile.mkdtemp(prefix="ddf-test-")
os.environ["CLAUDE_PLUGIN_DATA"] = TMP
os.environ["CLAUDE_PLUGIN_OPTION_OUTPUT_FOLDER"] = str(Path(TMP) / "out")

from ddforge import cli, evidence, ledger, machine, preflight, register, store  # noqa: E402

STATUT = """Стаття 1. Загальні положення
Товариство є юридичною особою за законодавством України.
Стаття 2. Органи управління
Загальні збори учасників є вищим органом Товариства.
Стаття 3. Відчуження частки
Відчуження частки у статутному капіталі потребує згоди загальних зборів учасників.
"""
CONTRACT_1 = """1. Предмет договору
1.1. Постачальник передає товар, а Покупець оплачує його.
2. Зміна контролю
2.1. У разі зміни контролю над Покупцем Постачальник має право розірвати договір в односторонньому порядку.
3. Прикінцеві положення
3.1. Договір набирає чинності з моменту підписання.
"""
CONTRACT_2 = """1. Предмет договору
1.1. Орендодавець передає Орендареві у строкове платне користування нежитлове приміщення за адресою, зазначеною у додатку до цього договору, а Орендар зобов'язується сплачувати орендну плату.
2. Прикінцеві положення
2.1. Договір набирає чинності з моменту підписання.
"""
LAW = """Стаття 1. Загальні положення
Цей Закон регулює відносини.
Стаття 2. Права контрагента
Сторона має право на відмову від договору, якщо це встановлено договором.
"""

BRIEF = {
    "object": "ТОВ Тестове",
    "depth": "full",
    "modules": ["general"],
    "legislation": {"ref": "law-nb", "title": "Законодавство"},
    "blocks": [{
        "id": "corporate", "title": "Корпоративний", "checklist": "custom",
        "notebook": {"ref": "nb-docs", "title": "Документи"},
        "custom_items": [
            {"id": "C1", "title": "Відчуження частки", "question": "Чи обмежено відчуження частки?", "tier": "core",
             "scope": "unit", "doc_kinds": ["statut"]},
            {"id": "C2", "title": "Зміна контролю", "question": "Чи є право розірвання при зміні контролю?",
             "tier": "core", "scope": "whole_doc", "doc_kinds": ["contract_commercial", "contract_lease"],
             "probes": ["змін\\w*\\s+контролю"]},
            {"id": "C3", "title": "Дозволи", "question": "Чи є дозволи?", "tier": "core", "scope": "unit",
             "doc_kinds": ["permit"]},
            {"id": "C4", "title": "Арбітраж", "question": "Чи є арбітражне застереження?", "tier": "extended",
             "scope": "whole_doc", "doc_kinds": ["contract_commercial", "contract_lease"], "probes": ["арбітраж"]},
        ],
    }],
}


def write_out(work, slot, payload):
    path = Path(work) / slot["output_file"]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def doc_by_source(work, source_id):
    return next(d for d in evidence.load_docs(work)["docs"].values() if d["source_id"] == source_id)


class MachineDryRun(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        brief_path = Path(TMP) / "brief.json"
        brief_path.write_text(json.dumps(BRIEF, ensure_ascii=False), encoding="utf-8")
        import io
        from contextlib import redirect_stdout
        buf = io.StringIO()
        with redirect_stdout(buf):
            cli.main(["task", "new", "--brief", str(brief_path)])
        info = json.loads(buf.getvalue())
        cls.work = Path(info["work"])
        (cls.work / "inbox" / "Статут ТОВ.txt").write_text(STATUT, encoding="utf-8")

    def slots(self):
        return list(store.read_json(self.work / "state.json")["wave"]["slots"].values())

    def test_full_run(self):
        work = self.work
        # 1. scouts
        action = machine.next_action(work)
        self.assertEqual(action["action"], "dispatch")
        self.assertEqual(len(action["slots"]), 2)
        for slot in self.slots():
            key = slot["meta"]["notebook_key"]
            sources = ([{"source_id": "L1", "title": "Цивільний кодекс України"}] if key == "law" else
                       [{"source_id": "s1", "title": "Договір поставки №1"},
                        {"source_id": "s2", "title": "Договір оренди №2"}])
            write_out(work, slot, {"notebook_key": key, "sources": sources, "blocker": None})
            self.assertEqual(machine.report_slot(work, slot["slot_id"])["status"], "done")
        # 2. gate
        action = machine.next_action(work)
        self.assertEqual(action["action"], "gate", action)
        self.assertIn("Позицій без документів", action["text"])
        self.assertEqual(machine.gate_reply(work, "proceed")["ok"], True)
        # 3. review: the statut is a project document, contracts are not fetched yet
        action = machine.next_action(work)
        self.assertEqual(action["action"], "dispatch", action)
        slot = self.slots()[0]
        self.assertEqual(slot["kind"], "review_batch")
        prompt = (work / slot["prompt_file"]).read_text(encoding="utf-8")
        self.assertIn("notebook_get", prompt)                      # nlm_protocol.md §8 block is delegated
        self.assertIn("Запасні джерела заборонені", prompt)         # documents notebook: no fallback
        self.assertIn("source_add заборонено", prompt)
        self.assertNotIn("source_list_drive`:", prompt.replace("(не source_list_drive)", ""))
        self.assertIn("Читати повністю", prompt)
        self.assertIn(str(work / slot["output_file"]), prompt)
        statut = next(d for d in evidence.load_docs(work)["docs"].values() if d["origin"] == "project")
        evidence.set_kind(work, statut["doc_id"], "statut")
        # the agent fetches and registers the contracts
        (Path(TMP) / "c1.txt").write_text(CONTRACT_1, encoding="utf-8")
        (Path(TMP) / "c2.txt").write_text(CONTRACT_2, encoding="utf-8")
        for sid, f in (("s1", "c1.txt"), ("s2", "c2.txt")):
            d = doc_by_source(work, sid)
            evidence.register_text(work, origin="notebook", notebook_key=d["notebook_key"], source_id=sid,
                                   title=d["title"], kind=d["kind"], file=str(Path(TMP) / f))
        c1, c2 = doc_by_source(work, "s1"), doc_by_source(work, "s2")
        self.assertEqual(c1["kind"], "contract_commercial")
        self.assertEqual(c2["kind"], "contract_lease")
        for unit in ("Стаття 1", "Стаття 2", "Стаття 3"):
            self.assertNotIn("error", evidence.cut(work, statut["doc_id"], unit=unit))
        scan2 = evidence.scan(work, [c1["doc_id"], c2["doc_id"]], ["змін\\w*\\s+контролю"], item_id="corporate:C2")
        evidence.cut(work, c1["doc_id"], scan_id=scan2["scan_id"], hit=1)
        scan4 = evidence.scan(work, [c1["doc_id"], c2["doc_id"]], ["арбітраж"], item_id="corporate:C4")
        self.assertEqual(sum(e["hits_total"] for e in scan4["docs"].values()), 0)
        quote = "У разі зміни контролю над Покупцем Постачальник має право розірвати договір в односторонньому порядку."
        write_out(work, slot, {"results": [
            {"item_id": "C1", "status": "clear", "basis": "units",
             "units_read": {statut["doc_id"]: ["ст.1", "ст.2", "ст.3"]}, "queries": 0},
            {"item_id": "C2", "status": "risk", "queries": 0, "findings": [{
                "doc_id": c1["doc_id"], "pinpoint": "п. 2.1", "fact": "Постачальник може розірвати договір при зміні контролю.",
                "issue": "Право розірвання при зміні контролю над Покупцем", "evidence": [{"cited_text": quote, "mode": "text"}],
                "severity_proposal": "critical", "consequence_hypothesis": "Втрата договору після закриття угоди",
                "seller_request": "Отримати згоду Постачальника"}]},
            {"item_id": "C3", "status": "clear", "basis": "chain", "queries": 1},
            {"item_id": "C4", "status": "clear", "basis": "scan", "queries": 0},
        ], "blocker": None})
        self.assertEqual(machine.report_slot(work, slot["slot_id"])["status"], "done")
        row = {r: ledger.load_ledger(work, "corporate")["items"][r] for r in ("C1", "C2", "C3", "C4")}
        self.assertEqual(row["C1"]["status"], "clear")
        self.assertEqual(row["C2"]["status"], "risk")
        self.assertEqual(row["C3"]["status"], "missing")   # no permit documents: pre-settled by the program
        self.assertEqual(row["C4"]["status"], "clear")
        # 4. legal basis
        action = machine.next_action(work)
        self.assertEqual(action["action"], "dispatch", action)
        slot = self.slots()[0]
        self.assertEqual(slot["kind"], "legal_basis")
        (Path(TMP) / "law.txt").write_text(LAW, encoding="utf-8")
        law = doc_by_source(work, "L1")
        evidence.register_text(work, origin="legislation", notebook_key="law", source_id="L1", title=law["title"],
                               kind="law", file=str(Path(TMP) / "law.txt"))
        cut = evidence.cut(work, law["doc_id"], unit="Стаття 2")
        fid = slot["meta"]["finding_ids"][0]
        write_out(work, slot, {"items": [{"finding_id": fid, "norm": {
            "source": "notebook", "doc_id": law["doc_id"], "act": "Цивільний кодекс України", "pinpoint": "ст. 2",
            "cited_text": "Сторона має право на відмову від договору, якщо це встановлено договором."},
            "consequence": "Контрагент може припинити договір"}]})
        self.assertEqual(machine.report_slot(work, slot["slot_id"])["status"], "done")
        self.assertEqual(register.findings(work)[0]["norm"]["verified"], "exact")
        # 5. calibration
        action = machine.next_action(work)
        slot = self.slots()[0]
        self.assertEqual(slot["kind"], "calibration")
        write_out(work, slot, {"items": [{"finding_id": fid, "severity": "critical", "rationale": "Втрата ключового договору.",
                                          "recommendation": {"type": "condition_precedent", "text": "Згода Постачальника до закриття."}}]})
        machine.report_slot(work, slot["slot_id"])
        # 6. audit
        action = machine.next_action(work)
        slot = self.slots()[0]
        self.assertEqual(slot["kind"], "audit")
        write_out(work, slot, {"verdicts": [{"block_id": b, "item_id": i, "ok": True, "reason": ""} for b, i in slot["meta"]["sample"]]})
        machine.report_slot(work, slot["slot_id"])
        # 7. findings gate (critical finding exists)
        action = machine.next_action(work)
        self.assertEqual(action["action"], "gate", action)
        self.assertEqual(action["gate_id"], "g-findings")
        machine.gate_reply(work, "proceed")
        # 8. drafting
        action = machine.next_action(work)
        slot = self.slots()[0]
        self.assertEqual(slot["kind"], "draft")
        bad = {"summary_points": ["Ризиків не виявлено."], "signing_questions": [], "block_summaries": {}}
        write_out(work, slot, bad)
        res = machine.report_slot(work, slot["slot_id"])
        self.assertEqual(res["status"], "open")          # rejected by the lint, retry offered
        good = {"summary_points": ["Виявлено одну критичну знахідку у договорі поставки."],
                "signing_questions": [{"text": "Чи отримано згоду Постачальника?", "finding_ids": [fid]}],
                "block_summaries": {"corporate": "Корпоративний блок: одна знахідка."}, "limitations_note": ""}
        write_out(work, slot, good)
        self.assertEqual(machine.report_slot(work, slot["slot_id"])["status"], "done")
        # 9. review
        action = machine.next_action(work)
        self.assertEqual(len(self.slots()), 2)
        for s in self.slots():
            write_out(work, s, {"issues": []})
            machine.report_slot(work, s["slot_id"])
        # 10. export
        action = machine.next_action(work)
        self.assertEqual(action["action"], "done", action)
        names = " ".join(action["outputs"])
        self.assertIn(".docx", names)
        self.assertIn(".xlsx", names)
        report_md = (Path(self.work) / "outputs").glob("*.md")
        text = next(report_md).read_text(encoding="utf-8")
        self.assertIn("Звіт не містить рішення про укладення угоди", text)
        self.assertIn("Не виявлено", text)
        self.assertIn("Документ відсутній", text)
        import docx
        import openpyxl
        docx_path = next(Path(p) for p in action["outputs"] if p.endswith(".docx") and "Звіт" in p)
        document = docx.Document(str(docx_path))
        self.assertTrue(any("Звіт про юридичну перевірку" in p.text for p in document.paragraphs))
        self.assertGreaterEqual(len(document.tables), 4)
        xlsx_path = next(Path(p) for p in action["outputs"] if p.endswith(".xlsx"))
        book = openpyxl.load_workbook(str(xlsx_path))
        self.assertEqual(book.sheetnames, ["Реєстр", "Покриття", "Запит документів", "Документи"])
        self.assertEqual(book["Реєстр"].max_row, 2)       # header + one finding
        self.assertTrue(any(p.endswith(".docx") and "Запит" in p for p in action["outputs"]))


class PauseAndDegrade(unittest.TestCase):
    def setUp(self):
        brief_path = Path(TMP) / "brief2.json"
        brief_path.write_text(json.dumps(BRIEF, ensure_ascii=False), encoding="utf-8")
        import io
        import time
        from contextlib import redirect_stdout
        time.sleep(1.1)   # task ids are second-resolution
        buf = io.StringIO()
        with redirect_stdout(buf):
            cli.main(["task", "new", "--brief", str(brief_path)])
        self.work = Path(json.loads(buf.getvalue())["work"])

    def slots(self):
        return list(store.read_json(self.work / "state.json")["wave"]["slots"].values())

    def test_auth_pause_resume_and_failed_slot(self):
        work = self.work
        machine.next_action(work)
        first, second = self.slots()
        write_out(work, first, {"notebook_key": first["meta"]["notebook_key"], "sources": [], "blocker": "auth_expired"})
        self.assertEqual(machine.report_slot(work, first["slot_id"])["status"], "paused")
        action = machine.next_action(work)
        self.assertEqual(action["action"], "pause")
        self.assertIn("nlm login", action["message"])
        machine.resume(work)
        action = machine.next_action(work)
        self.assertEqual(action["action"], "dispatch")
        open_ids = [s["slot_id"] for s in action["slots"]]
        self.assertIn(first["slot_id"], open_ids)
        # the second slot never returns valid JSON: three attempts, then degradation (the task goes on)
        for attempt in range(3):
            (work / second["output_file"]).write_text("не json", encoding="utf-8")
            result = machine.report_slot(work, second["slot_id"])
        self.assertEqual(result["status"], "failed")
        state = store.read_json(work / "state.json")
        self.assertTrue(any("не отримано" in d for d in state["degraded"]))
        retry_prompt = (work / second["prompt_file"]).read_text(encoding="utf-8")
        self.assertIn("Повторна спроба", retry_prompt)

    def test_cancel(self):
        machine.next_action(self.work)
        machine.cancel(self.work)
        self.assertEqual(machine.next_action(self.work)["action"], "cancelled")


if __name__ == "__main__":
    try:
        unittest.main()
    finally:
        if not os.environ.get("DDF_KEEP"):
            shutil.rmtree(TMP, ignore_errors=True)
        else:
            print("kept:", TMP)
