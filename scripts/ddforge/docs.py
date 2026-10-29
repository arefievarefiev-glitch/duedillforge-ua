"""Document inventory helpers: which documents belong to a block, and ingestion of project documents."""

from __future__ import annotations

from pathlib import Path

from . import evidence, store
from .vocab import UNDETERMINED

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".html", ".htm"}


def all_docs(work) -> list[dict]:
    return list(evidence.load_docs(work)["docs"].values())


def matches_filter(doc: dict, flt) -> bool:
    """`source_filter` of a block: None/[] = all; otherwise any entry equal to the source id or contained in the title."""
    if not flt:
        return True
    if isinstance(flt, dict):
        entries = list(flt.get("source_ids", [])) + list(flt.get("title_contains", []))
    else:
        entries = list(flt)
    title = (doc.get("title") or "").lower()
    for entry in entries:
        entry_l = str(entry).lower()
        if entry_l == str(doc.get("source_id", "")).lower() or entry_l in title:
            return True
    return False


def block_docs(work, state: dict, block_id: str) -> list[dict]:
    block = state["blocks"][block_id]
    key = block["notebook_key"]
    result = []
    for doc in all_docs(work):
        if doc["origin"] == "notebook" and doc.get("notebook_key") == key and matches_filter(doc, block.get("source_filter")):
            result.append(doc)
        elif doc["origin"] == "project":
            result.append(doc)
    return result


def law_docs(work) -> list[dict]:
    return [doc for doc in all_docs(work) if doc["origin"] == "legislation"]


# --- project documents ------------------------------------------------------------------------------


def _extract_docx(path: Path) -> str | None:
    try:
        import docx  # type: ignore
    except ImportError:
        return None
    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(path: Path) -> str | None:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        return None
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def ingest_inbox(work) -> dict:
    """Register every file in `<work>/inbox/` as a project document (kind: undetermined until classified)."""
    work = Path(work)
    inbox = work / "inbox"
    results = {"registered": [], "no_text": [], "unchanged": []}
    if not inbox.is_dir():
        return {**results, "note": "inbox is missing or empty"}
    registry = evidence.load_docs(work)["docs"]
    for path in sorted(p for p in inbox.rglob("*") if p.is_file()):
        rel = path.relative_to(inbox).as_posix()
        suffix = path.suffix.lower()
        text = None
        reason = "no_full_text"
        if suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
        elif suffix == ".docx":
            text = _extract_docx(path)
            if text is None:
                reason = "python-docx is not installed"
        elif suffix == ".pdf":
            text = _extract_pdf(path)
            if text is None:
                reason = "pypdf is not installed"
        known = next((d for d in registry.values() if d["origin"] == "project" and d["source_id"] == rel), None)
        if known and text is not None and known.get("text_sha") == store.sha256_text(evidence._clean(text)):
            results["unchanged"].append(known["doc_id"])
            continue
        if text is None:
            info = evidence.register_no_text(work, origin="project", notebook_key=None, source_id=rel,
                                             title=path.stem, kind=UNDETERMINED, reason=reason)
            results["no_text"].append({"doc_id": info["doc_id"], "file": rel, "reason": reason})
        else:
            info = evidence.register_text(work, origin="project", notebook_key=None, source_id=rel,
                                          title=path.stem, kind=UNDETERMINED, text=text)
            results["registered"].append({"doc_id": info["doc_id"], "file": rel,
                                          "text_quality": info["text_quality"], "chars": info["chars"]})
    return results
