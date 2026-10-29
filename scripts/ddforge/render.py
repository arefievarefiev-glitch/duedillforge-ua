"""Report rendering: one intermediate list of blocks -> markdown, DOCX; register -> XLSX; request list -> DOCX.

Status and severity wording, legends and fixed sentences come from `vocab.py`. Text written by a model enters only
through the writer's slots (summary, signing questions, block summaries) and is linted first.
"""

from __future__ import annotations

import csv
import re
import shutil
from pathlib import Path

from . import config, docs as docs_mod, evidence, ledger, register, store
from .vocab import (
    BASIS_UK, CONFIRMATION_UK, DEPTH_UK, DOC_KINDS_UK, INTERTEMPORAL_NOTE_UK, LDH_MARKER_UK,
    LEGISLATION_RESPONSIBILITY_UK, NO_RELIANCE_UK, OUTCOME_PHRASES, QUALITY_UK, REASONS_UK, REC_UK,
    REDFLAG_CLOSING_UK, RISK, SEVERITY_DEF_UK, SEVERITY_UK, STATUS_LEGEND_UK, STATUS_UK, CRITICAL,
)

FILL = {"critical": "F4CCCC", "significant": "FCE5CD", "info": "E7EDF3", "head": "D9D9D9"}
ID_RE = re.compile(r"\bR-[A-Z0-9]+-\d{2}\b")


# --- input for the writer and its lint ----------------------------------------------------------------------


def _norm_text(f: dict) -> str:
    norm = f.get("norm") or {}
    source = norm.get("source", "none")
    if source == "notebook":
        text = " ".join(x for x in (norm.get("act"), norm.get("pinpoint")) if x) or "норма з блокнота законодавства"
        if norm.get("verified") not in ("exact", "fuzzy"):
            text += " (цитату норми звірити)"
        return text
    if source == "ldh":
        text = " ".join(x for x in (norm.get("act"), norm.get("pinpoint")) if x) or "норма за Legal Data Hunter"
        return f"{text} {LDH_MARKER_UK}"
    need = norm.get("need") or ""
    return "Правову основу не встановлено" + (f": потрібна {need}" if need else "")


def _finding_view(f: dict) -> dict:
    rec = f.get("recommendation") or {}
    return {
        "id": f["id"], "block": f["block_id"], "severity": register.final_severity(f),
        "severity_uk": SEVERITY_UK[register.final_severity(f)], "document": f["doc_title"], "pinpoint": f["pinpoint"],
        "fact": f["fact"], "issue": f["issue"], "consequence": f.get("consequence") or f.get("consequence_hypothesis") or "",
        "norm": _norm_text(f), "confirmation": f["confirmation"],
        "recommendation_type": REC_UK.get(rec.get("type"), ""), "recommendation": rec.get("text", ""),
        "seller_request": f.get("seller_request", ""),
    }


def writer_input(work, state: dict) -> dict:
    work = Path(work)
    overall = ledger.overall(work, state)
    findings = [_finding_view(f) for f in register.sorted_findings(work)]
    blocks = []
    for bid in state["block_order"]:
        st = ledger.stats(ledger.load_ledger(work, bid))
        blocks.append({"id": bid, "title": state["blocks"][bid]["title"], "counts": st["counts"],
                       "confirmed_pct": st["confirmed_pct"], "screened_pct": st["screened_pct"],
                       "finding_ids": [f["id"] for f in findings if f["block"] == bid]})
    return {
        "object": state["object"], "buyer": state.get("buyer"), "transaction": state["transaction"],
        "depth": state["depth"], "depth_uk": DEPTH_UK[state["depth"]], "modules": state["modules"],
        "context": state.get("context", ""), "totals": overall["counts"], "confirmed_pct": overall["confirmed_pct"],
        "findings": findings, "blocks": blocks, "request_rows": len(_rfi(work, state)),
        "degraded": state.get("degraded", []),
        "output_contract": "див. стандартну інструкцію агента report-writer",
    }


def lint_writer(work, state: dict, data: dict) -> list[str]:
    problems = []
    points = data.get("summary_points")
    if not isinstance(points, list) or not points or not all(isinstance(p, str) and p.strip() for p in points):
        problems.append("summary_points: потрібен непорожній список рядків")
    if isinstance(points, list) and len(points) > 8:
        problems.append("summary_points: не більше 8 тез")
    questions = data.get("signing_questions")
    if not isinstance(questions, list):
        problems.append("signing_questions: потрібен список")
        questions = []
    summaries = data.get("block_summaries")
    if not isinstance(summaries, dict):
        problems.append("block_summaries: потрібен об'єкт")
        summaries = {}
    known = {f["id"]: f for f in register.findings(work)}
    texts = list(points) if isinstance(points, list) else []
    referenced = set()
    for q in questions:
        if not isinstance(q, dict) or not isinstance(q.get("text"), str) or not q["text"].strip():
            problems.append("signing_questions: кожен елемент має містити text")
            continue
        texts.append(q["text"])
        for fid in q.get("finding_ids") or []:
            if fid not in known:
                problems.append(f"signing_questions: невідомий finding_id {fid}")
            referenced.add(fid)
    for bid, text in summaries.items():
        if bid not in state["blocks"]:
            problems.append(f"block_summaries: невідомий блок {bid}")
        if isinstance(text, str):
            texts.append(text)
    note = data.get("limitations_note")
    if isinstance(note, str):
        texts.append(note)
    for text in texts:
        lowered = text.lower()
        for phrase in OUTCOME_PHRASES:
            if phrase in lowered:
                problems.append(f"текст не має заявляти про відсутність ризиків (знайдено «{phrase}»): статуси друкує програма")
                break
        for fid in ID_RE.findall(text):
            if fid not in known:
                problems.append(f"у тексті згадано неіснуючу знахідку {fid}")
        if len(text) > 1500:
            problems.append("рядок довший за 1500 знаків; скороти")
    critical = {f["id"] for f in known.values() if register.final_severity(f) == CRITICAL}
    if critical - referenced:
        problems.append("критичні знахідки без питання до підписання: " + ", ".join(sorted(critical - referenced)))
    return list(dict.fromkeys(problems))


def fallback_writer_text(work, state: dict) -> dict:
    findings = register.sorted_findings(work)
    by = {}
    for f in findings:
        by[register.final_severity(f)] = by.get(register.final_severity(f), 0) + 1
    points = [f"Знахідок усього: {len(findings)} (критичних — {by.get('critical', 0)}, суттєвих — "
              f"{by.get('significant', 0)}, до відома — {by.get('info', 0)})."]
    questions = [{"text": f"Вирішити питання за знахідкою {f['id']}: {f['issue']}", "finding_ids": [f["id"]]}
                 for f in findings if register.final_severity(f) == CRITICAL]
    return {"summary_points": points, "signing_questions": questions, "block_summaries": {},
            "limitations_note": "Текст резюме сформовано програмою зі структурованих даних реєстру."}


# --- helpers over ledgers ------------------------------------------------------------------------------------


def _rfi(work, state: dict) -> list[dict]:
    ledgers = ledger.all_ledgers(work, state)
    items_by_block = {bid: ledger.block_items(work, bid) for bid in state["block_order"]}
    return register.build_rfi(work, state, ledgers, items_by_block)


def _row_note(row: dict) -> str:
    st = row["status"]
    if st == "clear":
        text = BASIS_UK.get(row.get("basis"), "підтверджено програмою")
        parts = []
        for cov in (row.get("coverage") or {}).values():
            if cov.get("basis") in ("units", "chain"):
                parts.append(f"{cov['units_read']} з {cov['units_total']} одиниць")
            elif cov.get("basis") == "scan":
                parts.append(f"збігів переглянуто: {cov['hits']}")
        if parts:
            text += " (" + "; ".join(parts[:6]) + ")"
        return text
    if st == "unverified":
        text = REASONS_UK.get(row.get("reason"), "")
        if row.get("downgraded_from") == "clear":
            text += " (заявку «Не виявлено» програма не підтвердила)"
        return (text + (f": {row['detail']}" if row.get("detail") else "")).strip()
    if st == "na":
        return row.get("na_reason") or ""
    if st == "missing":
        kinds = "; ".join(DOC_KINDS_UK.get(m["kind"], m["kind"]) + (f" ({m['what']})" if m.get("what") else "")
                          for m in row.get("missing", []))
        return f"потрібні: {kinds}" if kinds else ""
    if st == "risk":
        return "знахідки: " + ", ".join(row.get("finding_ids", []))
    return ""


# --- report blocks ---------------------------------------------------------------------------------------------


def report_blocks(work, state: dict, writer: dict | None) -> list[dict]:
    work = Path(work)
    writer = writer or fallback_writer_text(work, state)
    overall = ledger.overall(work, state)
    findings = register.sorted_findings(work)
    rfi = _rfi(work, state)
    out = [{"t": "title", "text": "Звіт про юридичну перевірку (аудит)"},
           {"t": "p", "text": f"Об'єкт: {state['object']}"},
           {"t": "p", "text": f"Угода: {state['transaction']}" + (f". Покупець: {state['buyer']}" if state.get("buyer") else "")},
           {"t": "p", "text": f"Дата складання: {store.now_iso()[:10]}. Режим: {DEPTH_UK[state['depth']]}."},
           {"t": "h1", "text": "1. Резюме"}]
    out.append({"t": "ul", "items": list(writer["summary_points"])})
    counts = {}
    for f in findings:
        counts[register.final_severity(f)] = counts.get(register.final_severity(f), 0) + 1
    out.append({"t": "table", "header": ["Рівень", "Кількість", "Значення"],
                "rows": [[{"text": SEVERITY_UK[s], "fill": FILL[s]}, str(counts.get(s, 0)), SEVERITY_DEF_UK[s]]
                         for s in ("critical", "significant", "info")], "widths": [3, 2, 11]})
    out.append({"t": "h2", "text": "Питання, які потрібно вирішити до підписання"})
    if writer["signing_questions"]:
        out.append({"t": "ol", "items": [q["text"] + (" (" + ", ".join(q.get("finding_ids") or []) + ")"
                                                      if q.get("finding_ids") else "") for q in writer["signing_questions"]]})
    else:
        out.append({"t": "p", "text": "Питань, які потрібно вирішити до підписання, за результатами цієї перевірки не сформовано."})
    out.append({"t": "p", "text": "Звіт не містить рішення про укладення угоди; рішення належить замовнику."})
    if state["depth"] == "redflag":
        out.append({"t": "note", "text": REDFLAG_CLOSING_UK})
    # scope and limits
    totals = overall["counts"]
    out += [{"t": "h1", "text": "2. Обсяг перевірки і обмеження"},
            {"t": "p", "text": f"Позицій перевірки: {overall['total']}; блоків: {len(state['block_order'])}. "
                               f"Частка позицій із підтвердженим результатом (статуси «Виявлено», «Не виявлено», "
                               f"«Не застосовується»): {overall['confirmed_pct']}%."},
            {"t": "table", "header": ["Статус", "Позицій", "Що це означає"],
             "rows": [[STATUS_UK[s], str(totals.get(s, 0)), STATUS_LEGEND_UK[s]] for s in
                      ("risk", "clear", "screened", "unverified", "missing", "na") if s in STATUS_LEGEND_UK],
             "widths": [3, 2, 11]},
            {"t": "h2", "text": "Обмеження"}]
    limits = [LEGISLATION_RESPONSIBILITY_UK, INTERTEMPORAL_NOTE_UK,
              "Позначка «Не виявлено» ставиться лише коли програма підтвердила повноту перевірки; "
              "позначка «Скринінг без ознак» таким підтвердженням не є.", NO_RELIANCE_UK]
    limits += list(state.get("degraded", []))
    if writer.get("limitations_note"):
        limits.append(writer["limitations_note"])
    out.append({"t": "ul", "items": limits})
    # register
    out.append({"t": "h1", "text": "3. Реєстр знахідок"})
    if findings:
        rows = []
        for f in findings:
            v = _finding_view(f)
            rows.append([f["id"], {"text": v["severity_uk"], "fill": FILL[v["severity"]]},
                         state["blocks"][f["block_id"]]["title"], f"{v['document']}, {v['pinpoint']}", v["issue"],
                         (v["recommendation_type"] + (": " + v["recommendation"] if v["recommendation"] else ""))])
        out.append({"t": "table", "header": ["ID", "Рівень", "Блок", "Документ і положення", "Суть", "Рекомендація"],
                    "rows": rows, "widths": [2, 2, 2, 4, 5, 5]})
    else:
        out.append({"t": "p", "text": "Знахідок у реєстрі немає. Це не підтвердження відсутності ризиків: "
                                      "обсяг підтвердженої перевірки наведено в розділі 2."})
    # blocks
    for n, bid in enumerate(state["block_order"], start=4):
        block = state["blocks"][bid]
        led = ledger.load_ledger(work, bid)
        st = ledger.stats(led)
        out.append({"t": "h1", "text": f"{n}. {block['title']}"})
        summary = (writer.get("block_summaries") or {}).get(bid)
        if summary:
            out.append({"t": "p", "text": summary})
        out.append({"t": "p", "text": f"Позицій: {st['total']}. " + "; ".join(
            f"{STATUS_UK[s]}: {c}" for s, c in st["counts"].items() if c)})
        for f in [f for f in findings if f["block_id"] == bid]:
            v = _finding_view(f)
            out.append({"t": "h3", "text": f"{f['id']} · {v['severity_uk']} · {f['issue']}"})
            out.append({"t": "p", "text": f"Документ: {v['document']}. Положення: {v['pinpoint']}."})
            out.append({"t": "p", "text": f"Факт: {v['fact']}"})
            for ev in f.get("evidence", []):
                out.append({"t": "quote", "text": ev["cited_text"]})
            out.append({"t": "p", "text": f"Звірка цитати: {CONFIRMATION_UK[f['confirmation']]}."})
            if f.get("text_quality") != "ok":
                out.append({"t": "note", "text": f"Текст документа: {QUALITY_UK[f['text_quality']]}; звірити з оригіналом."})
            out.append({"t": "p", "text": f"Правова основа: {v['norm']}."})
            if v["consequence"]:
                out.append({"t": "p", "text": f"Можливий наслідок: {v['consequence']}"})
            if f.get("rationale"):
                out.append({"t": "p", "text": f"Обґрунтування рівня: {f['rationale']}"})
            if v["recommendation"] or v["recommendation_type"]:
                out.append({"t": "p", "text": f"Рекомендація ({v['recommendation_type']}): {v['recommendation']}"})
            if f.get("calibrated") is False:
                out.append({"t": "note", "text": "Рівень не пройшов калібрування: наведено пропозицію перевірника."})
            if f.get("intertemporal"):
                out.append({"t": "note", "text": INTERTEMPORAL_NOTE_UK})
            if f.get("user_note"):
                out.append({"t": "p", "text": f"Примітка юриста: {f['user_note']}"})
        items = {i["id"]: i for i in ledger.block_items(work, bid)}
        rows = []
        for iid, item in items.items():
            row = led["items"].get(iid)
            if not row:
                continue
            rows.append([iid, item["title"], STATUS_UK[row["status"]], _row_note(row)])
        out.append({"t": "h2", "text": "Покриття позицій блоку"})
        out.append({"t": "table", "header": ["Код", "Позиція", "Статус", "Підстава або причина"], "rows": rows,
                    "widths": [2, 5, 3, 7]})
    # request list
    n = 4 + len(state["block_order"])
    out.append({"t": "h1", "text": f"{n}. Запит додаткових документів"})
    if rfi:
        out.append({"t": "table", "header": ["№", "Що потрібно", "Чому", "Пріоритет"],
                    "rows": [[r["id"], r["what"], r["why"], r["priority"]] for r in rfi], "widths": [2, 8, 6, 2]})
    else:
        out.append({"t": "p", "text": "Додаткових документів за результатами перевірки не запитано."})
    # appendices
    out.append({"t": "h1", "text": f"{n + 1}. Додатки"})
    out.append({"t": "h2", "text": "Додаток А. Перелік документів, зареєстрованих у перевірці"})
    drows = [[d["doc_id"], d["title"], DOC_KINDS_UK.get(d.get("kind"), d.get("kind")),
              "не завантажено" if d.get("fetch_state") == "not_fetched" else QUALITY_UK[d["text_quality"]]]
             for d in docs_mod.all_docs(work) if d["origin"] != "legislation"]
    out.append({"t": "table", "header": ["doc_id", "Назва", "Вид", "Текст"], "rows": drows, "widths": [2, 8, 5, 3]})
    return out


def to_markdown(blocks: list[dict]) -> str:
    def cell(c):
        return (c["text"] if isinstance(c, dict) else str(c)).replace("|", "\\|").replace("\n", " ")

    lines = []
    for b in blocks:
        t = b["t"]
        if t == "title":
            lines += [f"# {b['text']}", ""]
        elif t in ("h1", "h2", "h3"):
            lines += [("#" * {"h1": 2, "h2": 3, "h3": 4}[t]) + " " + b["text"], ""]
        elif t == "p":
            lines += [b["text"], ""]
        elif t == "note":
            lines += [f"*{b['text']}*", ""]
        elif t == "quote":
            lines += ["> " + b["text"].replace("\n", "\n> "), ""]
        elif t == "ul":
            lines += [f"- {i}" for i in b["items"]] + [""]
        elif t == "ol":
            lines += [f"{n}. {i}" for n, i in enumerate(b["items"], 1)] + [""]
        elif t == "table":
            lines.append("| " + " | ".join(b["header"]) + " |")
            lines.append("|" + "---|" * len(b["header"]))
            for row in b["rows"]:
                lines.append("| " + " | ".join(cell(c) for c in row) + " |")
            lines.append("")
    return "\n".join(lines)


def draft_markdown(work, state: dict, writer: dict | None) -> str:
    return to_markdown(report_blocks(work, state, writer))


# --- DOCX ----------------------------------------------------------------------------------------------------------


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def _shade(cell, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def blocks_to_docx(blocks: list[dict], path: Path, footer: str = "") -> None:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor

    document = docx.Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11), ("Title", 18)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    if footer:
        section.footer.paragraphs[0].text = footer
    for b in blocks:
        t = b["t"]
        if t == "title":
            document.add_paragraph(b["text"], style="Title")
        elif t in ("h1", "h2", "h3"):
            document.add_heading(b["text"], level={"h1": 1, "h2": 2, "h3": 3}[t])
        elif t == "p":
            document.add_paragraph(b["text"])
        elif t == "note":
            p = document.add_paragraph()
            run = p.add_run(b["text"])
            run.italic = True
        elif t == "quote":
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run("«" + b["text"] + "»")
            run.italic = True
        elif t == "ul":
            for item in b["items"]:
                document.add_paragraph(item, style="List Bullet")
        elif t == "ol":
            for item in b["items"]:
                document.add_paragraph(item, style="List Number")
        elif t == "table":
            table = document.add_table(rows=1, cols=len(b["header"]))
            table.style = "Table Grid"
            widths = b.get("widths") or [1] * len(b["header"])
            total = sum(widths)
            usable = 17.0
            for i, head in enumerate(b["header"]):
                c = table.rows[0].cells[i]
                c.text = ""
                run = c.paragraphs[0].add_run(head)
                run.bold = True
                run.font.size = Pt(9.5)
                _shade(c, FILL["head"])
            for row in b["rows"]:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    text = value["text"] if isinstance(value, dict) else str(value)
                    cells[i].text = ""
                    run = cells[i].paragraphs[0].add_run(text)
                    run.font.size = Pt(9.5)
                    if isinstance(value, dict) and value.get("fill"):
                        _shade(cells[i], value["fill"])
            for row in table.rows:
                for i, c in enumerate(row.cells):
                    c.width = Cm(usable * widths[i] / total)
            document.add_paragraph()
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def rfi_blocks(work, state: dict) -> list[dict]:
    rfi = _rfi(work, state)
    out = [{"t": "title", "text": "Запит додаткових документів та інформації"},
           {"t": "p", "text": f"Об'єкт перевірки: {state['object']}. Дата: {store.now_iso()[:10]}."},
           {"t": "p", "text": "Для завершення перевірки просимо надати документи і відомості, перелічені нижче. "
                              "Пріоритет «високий» означає, що без документа не можна завершити перевірку ключових позицій."}]
    if rfi:
        out.append({"t": "table", "header": ["№", "Документ або відомості", "Чому потрібні", "Пріоритет"],
                    "rows": [[r["id"], r["what"], r["why"], r["priority"]] for r in rfi], "widths": [2, 8, 6, 2]})
    else:
        out.append({"t": "p", "text": "Додаткових документів не запитано."})
    return out


# --- XLSX ------------------------------------------------------------------------------------------------------------


def write_xlsx(work, state: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def sheet(title, header, rows, widths, first=False, fill_col=None):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor=FILL["head"])
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows:
            ws.append([r["text"] if isinstance(r, dict) else r for r in row])
            if fill_col is not None and isinstance(row[fill_col], dict) and row[fill_col].get("fill"):
                ws.cell(row=ws.max_row, column=fill_col + 1).fill = PatternFill("solid", fgColor=row[fill_col]["fill"])
            for c in ws[ws.max_row]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    findings = register.sorted_findings(work)
    rows = []
    for f in findings:
        v = _finding_view(f)
        rows.append([f["id"], {"text": v["severity_uk"], "fill": FILL[v["severity"]]}, f["block_id"], v["document"],
                     v["pinpoint"], v["fact"], v["issue"], v["norm"], v["consequence"], v["recommendation_type"],
                     v["recommendation"], v["seller_request"], CONFIRMATION_UK[f["confirmation"]]])
    sheet("Реєстр", ["ID", "Рівень", "Блок", "Документ", "Положення", "Факт", "Суть", "Правова основа", "Наслідок",
                     "Тип рекомендації", "Рекомендація", "Запит до продавця", "Звірка цитати"], rows,
          [9, 12, 12, 30, 18, 40, 40, 30, 35, 22, 40, 35, 30], first=True, fill_col=1)
    rows = []
    for bid in state["block_order"]:
        led = ledger.load_ledger(work, bid)
        for item in ledger.block_items(work, bid):
            row = led["items"].get(item["id"])
            if row:
                rows.append([state["blocks"][bid]["title"], item["id"], item["title"], item["tier"],
                             STATUS_UK[row["status"]], _row_note(row), len(row.get("scope_docs", []))])
    sheet("Покриття", ["Блок", "Код", "Позиція", "Рівень позиції", "Статус", "Підстава або причина", "Документів"],
          rows, [20, 10, 45, 12, 22, 60, 12])
    rfi = _rfi(work, state)
    sheet("Запит документів", ["№", "Що потрібно", "Чому", "Пріоритет", "Блоки"],
          [[r["id"], r["what"], r["why"], r["priority"], ", ".join(r["blocks"])] for r in rfi], [8, 60, 45, 12, 25])
    sheet("Документи", ["doc_id", "Назва", "Вид", "Текст", "Знаків", "Одиниць"],
          [[d["doc_id"], d["title"], DOC_KINDS_UK.get(d.get("kind"), d.get("kind")),
            "не завантажено" if d.get("fetch_state") == "not_fetched" else QUALITY_UK[d["text_quality"]],
            d.get("chars", 0), d.get("units", 0)] for d in docs_mod.all_docs(work) if d["origin"] != "legislation"],
          [10, 60, 40, 22, 10, 10])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def write_csv_fallback(work, state: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["ID", "Рівень", "Блок", "Документ", "Положення", "Суть", "Рекомендація"])
        for f in register.sorted_findings(work):
            v = _finding_view(f)
            writer.writerow([f["id"], v["severity_uk"], f["block_id"], v["document"], v["pinpoint"], v["issue"],
                             v["recommendation"]])


# --- export ----------------------------------------------------------------------------------------------------------


def export(work, state: dict, writer: dict | None) -> dict:
    work = Path(work)
    slug = state["slug"]
    out_dir = work / "outputs"
    out_dir.mkdir(parents=True, exist_ok=True)
    warnings, files = [], []
    blocks = report_blocks(work, state, writer)
    footer = f"{state['object']} · звіт про юридичну перевірку · {store.now_iso()[:10]}"
    md_path = out_dir / f"Звіт_перевірки_{slug}.md"
    store.write_text(md_path, to_markdown(blocks))
    if _docx_available():
        docx_path = out_dir / f"Звіт_перевірки_{slug}.docx"
        blocks_to_docx(blocks, docx_path, footer)
        files.append(docx_path)
        rfi_path = out_dir / f"Запит_документів_{slug}.docx"
        blocks_to_docx(rfi_blocks(work, state), rfi_path, footer)
        files.append(rfi_path)
    else:
        warnings.append("python-docx недоступний: звіт видано у форматі markdown")
        files.append(md_path)
    try:
        xlsx_path = out_dir / f"Реєстр_ризиків_{slug}.xlsx"
        write_xlsx(work, state, xlsx_path)
        files.append(xlsx_path)
    except ImportError:
        csv_path = out_dir / f"Реєстр_ризиків_{slug}.csv"
        write_csv_fallback(work, state, csv_path)
        files.append(csv_path)
        warnings.append("openpyxl недоступний: реєстр видано у форматі CSV")
    if state.get("unresolved_review_issues"):
        warnings.append(f"лишилося обов'язкових зауважень рецензентів: {len(state['unresolved_review_issues'])}")
    final = [str(p) for p in files]
    for target in _targets(slug):
        try:
            target.mkdir(parents=True, exist_ok=True)
            for p in files:
                shutil.copy2(p, target / p.name)
            final.append(f"копію збережено: {target}")
        except OSError as exc:
            warnings.append(f"не вдалося скопіювати у {target}: {exc}")
    return {"files": final, "warnings": warnings}


def _targets(slug: str) -> list[Path]:
    result = []
    folder = config.resolve_output_folder(None)
    if folder:
        result.append(folder / slug)
    publish = config.resolve_publish_folder(None, slug)
    if publish and publish not in result:
        result.append(publish)
    return result
